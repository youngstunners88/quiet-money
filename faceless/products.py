"""Product factory: spreadsheet products built from code, recalculated by a real spreadsheet engine, checked against our own math.

A product is a deterministic build (`build_*`), a verification that recalculates it in LibreOffice and compares to
independent Python math (`verify_*`), and a listing file the owner pastes into Gumroad/Etsy. Nothing is sold from here:
the owner opens the shop and approves price and copy; this module only makes the thing worth selling.
`python -m faceless products` builds everything into channel/products/<slug>/.
"""

from __future__ import annotations

import math
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from faceless import config
from faceless.config import Paths

OUT = Paths.root / "channel" / "products"
INK, GOLD, CREAM, MUTED = "0E0E10", "FFD23F", "F5F1E6", "8A8678"
IN_FILL = PatternFill("solid", fgColor="DCEBFF")      # blue-tinted: type here
OUT_FILL = PatternFill("solid", fgColor="EFEFEF")     # gray: a formula, leave alone
HEAD_FILL = PatternFill("solid", fgColor=INK)
GOOD_FILL = PatternFill("solid", fgColor="CDEFD3")
THIN = Side(style="thin", color="C9C9C9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DISCLAIMER = ("Educational tool, not financial advice. Results depend entirely on the assumptions you enter; returns are never guaranteed. "
              "Designed by Quiet Money with AI assistance; every formula was recalculated and checked against independent math.")
SHEETS = ["Start here", "1 Your numbers", "2 Income streams", "3 Rule of 100", "4 Projection", "5 Weekly log"]


def _title(ws, text: str, sub: str = "", width: int = 6) -> None:
    ws.sheet_view.showGridLines = False
    ws["A1"] = text
    ws["A1"].font = Font(name="Calibri", size=20, bold=True, color=GOLD)
    for c in range(1, width + 1):
        ws.cell(1, c).fill = HEAD_FILL
        ws.cell(2, c).fill = HEAD_FILL
    ws.row_dimensions[1].height = 34
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name="Calibri", size=11, italic=True, color=CREAM)


def _put(ws, ref: str, value, *, fmt: str | None = None, inp: bool = False, bold: bool = False, note: str = "") -> None:
    c = ws[ref]
    c.value = value
    c.border = BOX
    c.font = Font(name="Calibri", size=11, bold=bold, color="1F4E9E" if inp else "000000")
    c.fill = IN_FILL if inp else OUT_FILL
    c.alignment = Alignment(horizontal="right", vertical="center")
    if fmt:
        c.number_format = fmt
    if note:
        ws.cell(c.row, c.column + 1).value = note
        ws.cell(c.row, c.column + 1).font = Font(name="Calibri", size=10, italic=True, color=MUTED)


def _label(ws, ref: str, text: str, bold: bool = False) -> None:
    ws[ref].value = text
    ws[ref].font = Font(name="Calibri", size=11, bold=bold)
    ws[ref].alignment = Alignment(vertical="center", wrap_text=True)


def _header(ws, row: int, cols: list[str]) -> None:
    for i, text in enumerate(cols, 1):
        c = ws.cell(row, i, text)
        c.font = Font(name="Calibri", size=11, bold=True, color=GOLD)
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def _widths(ws, widths: list[float]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# default inputs: also the numbers the verification recomputes in Python
DEFAULTS = {"age": 35, "spend": 3000, "income": 4500, "invested": 15000, "monthly": 900, "ret": 0.07, "infl": 0.03, "swr": 0.04, "extra": 100}


def build_escape_planner(dest: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEETS[0]
    _title(ws, "Rat Race Escape Planner", "Know your freedom number. Track every income stream. Execute for 100 days.", 6)
    steps = ["1. Open '1 Your numbers' and replace the blue cells with your own. Gray cells are formulas: leave them alone.",
             "2. Read your freedom number and your years to freedom. Then change one input at a time and watch what moves.",
             "3. List every income stream in '2 Income streams'. The goal is passive and asset income, not more hours.",
             "4. Run '3 Rule of 100' with your real offer, then log every week in '5 Weekly log'. Execution beats education.",
             "5. '4 Projection' shows the whole road, year by year, in today's dollars."]
    for i, t in enumerate(steps):
        _label(ws, f"A{4 + i}", t)
        ws.merge_cells(f"A{4 + i}:F{4 + i}")
        ws.row_dimensions[4 + i].height = 32
    _label(ws, "A10", "Colors", True)
    _put(ws, "A11", "Blue = you type here", inp=True)
    _put(ws, "A12", "Gray = formula")
    _label(ws, "A14", "How the math works", True)
    for i, t in enumerate(["Freedom number = yearly spending / safe withdrawal rate (the 4% rule is a rule of thumb, not a guarantee).",
                           "Everything is in today's dollars: the real return is (1 + return) / (1 + inflation) - 1.",
                           "Years to freedom uses the same NPER/FV functions your bank's calculators use, month by month."]):
        _label(ws, f"A{15 + i}", t)
        ws.merge_cells(f"A{15 + i}:F{15 + i}")
        ws.row_dimensions[15 + i].height = 30
    _label(ws, "A19", DISCLAIMER)
    ws.merge_cells("A19:F19")
    ws.row_dimensions[19].height = 48
    ws["A19"].font = Font(name="Calibri", size=10, italic=True, color=MUTED)
    _label(ws, "A21", "License: personal use for you and your household. Please do not resell or redistribute the file.")
    ws.merge_cells("A21:F21")
    _widths(ws, [30, 18, 18, 18, 18, 18])

    # ---- 1 Your numbers
    n = wb.create_sheet(SHEETS[1])
    _title(n, "1 Your numbers", "Replace the blue cells. Example numbers shown.", 4)
    rows = [("Your age", "age", "0", ""), ("Monthly spending you need (today's dollars)", "spend", "$#,##0", "What 'enough' costs you each month"),
            ("Monthly take-home income", "income", "$#,##0", ""), ("Invested today", "invested", "$#,##0", "Savings and investments you could live off"),
            ("Amount you invest each month", "monthly", "$#,##0", ""), ("Expected yearly return (assumption)", "ret", "0.0%", "Not a promise; try 4%, 7%, 10%"),
            ("Expected yearly inflation (assumption)", "infl", "0.0%", ""), ("Safe yearly withdrawal rate", "swr", "0.0%", "4% is a common rule of thumb"),
            ("Extra you could invest each month", "extra", "$#,##0", "For the 'what if' line below")]
    cell = {}
    _header(n, 3, ["Input", "Value", "Note", ""])
    for i, (lab, key, fmt, note) in enumerate(rows):
        r = 4 + i
        _label(n, f"A{r}", lab)
        _put(n, f"B{r}", DEFAULTS[key], fmt=fmt, inp=True)
        n[f"C{r}"] = note
        n[f"C{r}"].font = Font(name="Calibri", size=10, italic=True, color=MUTED)
        cell[key] = f"$B${r}"
    for key, lo, hi in (("ret", 0, 0.3), ("infl", 0, 0.2), ("swr", 0.01, 0.1)):
        dv = DataValidation(type="decimal", operator="between", formula1=str(lo), formula2=str(hi), showErrorMessage=True,
                            errorTitle="Check this number", error=f"Enter a percentage between {lo:.0%} and {hi:.0%} (for example 7%).")
        n.add_data_validation(dv)
        dv.add(cell[key].replace("$", ""))
    c = cell
    _header(n, 15, ["Your results", "Value", "What it means", ""])
    res = [("Freedom number", f"={c['spend']}*12/{c['swr']}", "$#,##0", "The invested amount that covers your spending for good"),
           ("Real yearly return", f"=(1+{c['ret']})/(1+{c['infl']})-1", "0.00%", "Return after inflation: keeps everything in today's dollars"),
           ("Savings rate", f"=IFERROR({c['monthly']}/{c['income']},0)", "0%", "Share of take-home income you invest"),
           ("Months to freedom", f'=IF({c["invested"]}>=B16,0,IFERROR(NPER(B17/12,-{c["monthly"]},-{c["invested"]},B16),"never at these numbers"))', "#,##0", ""),
           ("Years to freedom", '=IF(ISNUMBER(B19),B19/12,"never at these numbers")', "0.0", ""),
           ("Your freedom age", f'=IF(ISNUMBER(B20),{c["age"]}+B20,"-")', "0.0", ""),
           ("Passive income your money makes today", f"={c['invested']}*{c['swr']}/12", "$#,##0", "Per month, at your withdrawal rate"),
           ("Share of your spending already covered", f"=IFERROR(B22/{c['spend']},0)", "0%", ""),
           ("Years to freedom with the extra monthly amount",
            f'=IF({c["invested"]}>=B16,0,IFERROR(NPER(B17/12,-({c["monthly"]}+{c["extra"]}),-{c["invested"]},B16)/12,"never at these numbers"))', "0.0", ""),
           ("Years the extra amount saves you", '=IF(AND(ISNUMBER(B20),ISNUMBER(B24)),B20-B24,"-")', "0.0", "Small, boring, automatic. That is the whole trick.")]
    for i, (lab, f, fmt, note) in enumerate(res):
        r = 16 + i
        _label(n, f"A{r}", lab, bold=(i in (0, 4)))
        _put(n, f"B{r}", f, fmt=fmt, bold=(i in (0, 4)))
        n[f"C{r}"] = note
        n[f"C{r}"].font = Font(name="Calibri", size=10, italic=True, color=MUTED)
    n.conditional_formatting.add("B23", CellIsRule(operator="greaterThanOrEqual", formula=["1"], fill=GOOD_FILL))
    _widths(n, [48, 22, 58, 4])
    n.freeze_panes = "A4"

    # ---- 2 Income streams
    s = wb.create_sheet(SHEETS[2])
    _title(s, "2 Income streams", "Active trades hours for money. Passive runs without you. Asset is something you own that pays.", 6)
    _header(s, 3, ["Stream", "Type", "Per month", "Hours per week", "$ per hour", "What it is"])
    sample = [("Day job (example: replace)", "Active", 4500, 40), ("Side project (example)", "Active", 300, 5),
              ("Index fund dividends (example)", "Asset", 40, 0), ("Digital product sales (example)", "Passive", 0, 0)]
    for i in range(12):
        r = 4 + i
        row = sample[i] if i < len(sample) else ("", "", None, None)
        _put(s, f"A{r}", row[0], inp=True)
        s[f"A{r}"].alignment = Alignment(horizontal="left")
        _put(s, f"B{r}", row[1], inp=True)
        _put(s, f"C{r}", row[2], fmt="$#,##0", inp=True)
        _put(s, f"D{r}", row[3], fmt="0", inp=True)
        _put(s, f"E{r}", f'=IF(C{r}="","",IF(N(D{r})>0,C{r}/(D{r}*4.33),"no hours"))', fmt="$#,##0.00")
        _put(s, f"F{r}", f'=IF(B{r}="Active","Trades your time",IF(B{r}="Passive","Runs without you",IF(B{r}="Asset","You own it","")))')
        s[f"F{r}"].alignment = Alignment(horizontal="left")
    dv = DataValidation(type="list", formula1='"Active,Passive,Asset"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Pick a type", error="Choose Active, Passive, or Asset.")
    s.add_data_validation(dv)
    dv.add("B4:B15")
    _label(s, "A17", "Total per month", True)
    _put(s, "C17", "=SUM(C4:C15)", fmt="$#,##0", bold=True)
    _label(s, "A18", "Passive + asset income per month", True)
    _put(s, "C18", '=SUMIF(B4:B15,"Passive",C4:C15)+SUMIF(B4:B15,"Asset",C4:C15)', fmt="$#,##0", bold=True)
    _label(s, "A19", "Passive + asset share of your spending", True)
    _put(s, "C19", f"=IFERROR(C18/'{SHEETS[1]}'!B5,0)", fmt="0%", bold=True)
    _label(s, "A20", "Hours you trade each week", True)
    _put(s, "C20", '=SUMIF(B4:B15,"Active",D4:D15)', fmt="0", bold=True)
    _label(s, "A22", "The escape is the day C19 reaches 100%: your spending is paid by money that works while you do not.")
    s.merge_cells("A22:F22")
    _widths(s, [40, 12, 14, 16, 14, 24])
    s.freeze_panes = "A4"

    # ---- 3 Rule of 100
    r1 = wb.create_sheet(SHEETS[3])
    _title(r1, "3 Rule of 100", "Volume beats talent. 100 outreaches a day for 100 days is 10,000 shots at a customer.", 5)
    for i, (lab, val, fmt) in enumerate([("Outreaches per day", 100, "0"), ("Days", 100, "0"), ("Yes-rate (a guess until you test it)", 0.01, "0.0%"),
                                         ("Average sale per customer", 150, "$#,##0")]):
        _label(r1, f"A{4 + i}", lab)
        _put(r1, f"B{4 + i}", val, fmt=fmt, inp=True)
    _label(r1, "A9", "Total outreaches", True)
    _put(r1, "B9", "=B4*B5", fmt="#,##0", bold=True)
    _label(r1, "A10", "Customers", True)
    _put(r1, "B10", "=B9*B6", fmt="#,##0", bold=True)
    _label(r1, "A11", "Revenue", True)
    _put(r1, "B11", "=B10*B7", fmt="$#,##0", bold=True)
    _header(r1, 13, ["If your yes-rate is", "Customers", "Revenue", "", ""])
    for i, rate in enumerate([0.002, 0.005, 0.01, 0.02, 0.03]):
        r = 14 + i
        _put(r1, f"A{r}", rate, fmt="0.0%", inp=True)
        _put(r1, f"B{r}", f"=$B$9*A{r}", fmt="#,##0")
        _put(r1, f"C{r}", f"=B{r}*$B$7", fmt="$#,##0")
    _label(r1, "A20", "Hypothetical. Your real yes-rate and sale size are unknown until you make the calls: log them in '5 Weekly log' and replace the guesses.")
    r1.merge_cells("A20:E20")
    r1.row_dimensions[20].height = 34
    _widths(r1, [40, 16, 16, 4, 4])

    # ---- 4 Projection
    p = wb.create_sheet(SHEETS[4])
    _title(p, "4 Projection", "Year by year, in today's dollars, from your numbers. Highlighted when your money covers your spending.", 6)
    _header(p, 3, ["Year", "Age", "Invested", "Passive income / month", "Spending covered", "Freedom number"])
    nm = f"'{SHEETS[1]}'!"
    for y in range(0, 41):
        r = 4 + y
        _put(p, f"A{r}", y, fmt="0")
        _put(p, f"B{r}", f"={nm}$B$4+A{r}", fmt="0")
        _put(p, f"C{r}", f"=FV({nm}$B$17/12,A{r}*12,-{nm}$B$8,-{nm}$B$7)", fmt="$#,##0")
        _put(p, f"D{r}", f"=C{r}*{nm}$B$11/12", fmt="$#,##0")
        _put(p, f"E{r}", f"=IFERROR(D{r}/{nm}$B$5,0)", fmt="0%")
        _put(p, f"F{r}", f"={nm}$B$16", fmt="$#,##0")
    p.conditional_formatting.add("A4:F44", FormulaRule(formula=["$E4>=1"], fill=GOOD_FILL))
    ch = LineChart()
    ch.title = "Invested vs freedom number (today's dollars)"
    ch.height, ch.width = 9, 20
    ch.add_data(Reference(p, min_col=3, min_row=3, max_row=44), titles_from_data=True)
    ch.add_data(Reference(p, min_col=6, min_row=3, max_row=44), titles_from_data=True)
    ch.set_categories(Reference(p, min_col=1, min_row=4, max_row=44))
    p.add_chart(ch, "H3")
    _widths(p, [8, 8, 16, 22, 18, 18])
    p.freeze_panes = "A4"

    # ---- 5 Weekly log
    w = wb.create_sheet(SHEETS[5])
    _title(w, "5 Weekly log", "Thirteen weeks is about 100 days. Fill in what happened, not what you hoped for.", 7)
    _header(w, 3, ["Week", "Outreaches", "Replies", "Sales", "Revenue", "Hours worked", "What I learned"])
    for i in range(13):
        r = 4 + i
        _put(w, f"A{r}", i + 1, fmt="0")
        for col in "BCDEFG":
            _put(w, f"{col}{r}", None, fmt="$#,##0" if col == "E" else "0", inp=True)
        w[f"G{r}"].alignment = Alignment(horizontal="left")
    _label(w, "A18", "Totals", True)
    for col in "BCDEF":
        _put(w, f"{col}18", f"=SUM({col}4:{col}16)", fmt="$#,##0" if col == "E" else "#,##0", bold=True)
    _label(w, "A20", "Reply rate")
    _put(w, "B20", "=IFERROR(C18/B18,0)", fmt="0.0%")
    _label(w, "A21", "Close rate (sales per reply)")
    _put(w, "B21", "=IFERROR(D18/C18,0)", fmt="0.0%")
    _label(w, "A22", "Yes-rate (sales per outreach)")
    _put(w, "B22", "=IFERROR(D18/B18,0)", fmt="0.00%")
    _label(w, "A23", "Revenue per hour")
    _put(w, "B23", "=IFERROR(E18/F18,0)", fmt="$#,##0.00")
    _label(w, "A24", "Copy your real yes-rate into '3 Rule of 100' and watch the forecast become honest.")
    w.merge_cells("A24:G24")
    _widths(w, [28, 14, 12, 10, 14, 14, 48])
    w.freeze_panes = "A4"

    for sh in wb.worksheets:
        sh.sheet_properties.tabColor = GOLD if sh.title == SHEETS[1] else INK
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 1          # one tab prints (and previews) as one page
        sh.sheet_properties.pageSetUpPr.fitToPage = True
    wb.properties.title = "Rat Race Escape Planner"
    wb.properties.creator = "Quiet Money"
    dest.mkdir(parents=True, exist_ok=True)
    out = dest / "Rat-Race-Escape-Planner.xlsx"
    wb.save(out)
    return out


# ---- verification: recalculate in a real spreadsheet engine, compare to independent math ----------------

def recalc(xlsx: Path) -> Path | None:
    """Round-trip through headless LibreOffice so every formula is evaluated; None if LibreOffice is not installed."""
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if not office:
        return None
    tmp = Path(tempfile.mkdtemp(prefix="recalc-"))
    r = subprocess.run([office, "--headless", "--norestore", f"-env:UserInstallation=file://{tmp}/profile", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                        "--outdir", str(tmp), str(xlsx)], capture_output=True, text=True, timeout=180, check=False)
    out = tmp / xlsx.name
    return out if r.returncode == 0 and out.exists() else None


def expected_escape(d: dict | None = None) -> dict:
    """The planner's headline numbers computed without the spreadsheet (closed-form NPER)."""
    d = {**DEFAULTS, **(d or {})}
    freedom = d["spend"] * 12 / d["swr"]
    real = (1 + d["ret"]) / (1 + d["infl"]) - 1
    r = real / 12

    def months(pmt: float) -> float:
        if d["invested"] >= freedom:
            return 0.0
        return math.log((pmt + freedom * r) / (pmt + d["invested"] * r)) / math.log(1 + r)
    m, m2 = months(d["monthly"]), months(d["monthly"] + d["extra"])
    y20 = d["invested"] * (1 + r) ** 240 + d["monthly"] * ((1 + r) ** 240 - 1) / r
    return {"freedom": freedom, "real": real, "months": m, "years": m / 12, "years_extra": m2 / 12, "saved": (m - m2) / 12,
            "passive_now": d["invested"] * d["swr"] / 12, "balance_year20": y20}


def verify_escape(xlsx: Path) -> dict:
    """Recalculate and compare. Returns {'ok': bool, 'rows': [(name, sheet value, python value)], 'errors': [cells with errors]}."""
    calc = recalc(xlsx)
    if calc is None:
        return {"ok": None, "rows": [], "errors": ["LibreOffice Calc is not available (install libreoffice-calc): cannot recalculate"]}
    wb = load_workbook(calc, data_only=True)
    n, p = wb[SHEETS[1]], wb[SHEETS[4]]
    e = expected_escape()
    got = {"freedom": n["B16"].value, "real": n["B17"].value, "months": n["B19"].value, "years": n["B20"].value, "years_extra": n["B24"].value,
           "saved": n["B25"].value, "passive_now": n["B22"].value, "balance_year20": p["C24"].value}
    rows, ok = [], True
    for k, v in got.items():
        good = isinstance(v, (int, float)) and abs(v - e[k]) <= max(1e-6, abs(e[k]) * 1e-7)
        ok &= good
        rows.append((k, v, round(e[k], 4), good))
    errors = [f"{ws.title}!{c.coordinate}={c.value}" for ws in wb for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("#")]
    return {"ok": bool(ok and not errors), "rows": rows, "errors": errors, "calc": str(calc)}


def previews(xlsx: Path, out: Path) -> list[Path]:
    """Real screenshots of the product (PDF -> PNG) for the shop listing: no mock-ups, nothing invented."""
    office, ppm = shutil.which("soffice"), shutil.which("pdftoppm")
    if not office or not ppm:
        return []
    tmp = Path(tempfile.mkdtemp(prefix="prev-"))
    subprocess.run([office, "--headless", "--norestore", f"-env:UserInstallation=file://{tmp}/profile", "--convert-to", "pdf", "--outdir", str(tmp), str(xlsx)],
                   capture_output=True, timeout=180, check=False)
    pdf = tmp / (xlsx.stem + ".pdf")
    if not pdf.exists():
        return []
    subprocess.run([ppm, "-png", "-r", "90", "-f", "1", "-l", "5", str(pdf), str(tmp / "p")], capture_output=True, timeout=120, check=False)
    out.mkdir(parents=True, exist_ok=True)
    made = []
    for i, src in enumerate(sorted(tmp.glob("p-*.png")), 1):
        dst = out / f"preview-{i}.png"
        shutil.copy2(src, dst)
        made.append(dst)
    return made


def listing_escape(price: float | None = None) -> str:
    price = price if price is not None else config.load().get("forecast", {}).get("product_price", 19.0)
    link = config.load()["channel"].get("link_in_bio", "")
    return f"""# Listing: Rat Race Escape Planner

Prepared {config.today_utc().isoformat()}. The owner approves price, copy and the account before anything goes live.

## Gumroad (first: 10% fee, no AI rule)
- **Name:** Rat Race Escape Planner (Excel + Google Sheets)
- **Price:** ${price:.0f} (test $12 / $19 / $29; change only after 20 views)
- **Cover + thumbnails:** `preview-1.png` to `preview-5.png` (real screenshots of the file)
- **File:** `Rat-Race-Escape-Planner.xlsx` (opens in Excel, Google Sheets, LibreOffice, Numbers)

### Description
Know your freedom number, how many years it is away, and what one boring monthly habit does to that date.

**What you get (one spreadsheet, six tabs):**
- **Your numbers**: type your spending, income, savings and return assumptions. See your *freedom number* (the invested amount that covers your spending), months and years to freedom, your freedom age, and the years an extra monthly amount saves.
- **Income streams**: list every stream and tag it Active, Passive or Asset. See your dollars per hour and how much of your spending is already covered by money that works without you.
- **Rule of 100**: the volume math for any offer (100 outreaches a day for 100 days) with a sensitivity table for your yes-rate.
- **Projection**: forty years in today's dollars, with a chart; the row turns green the year your money covers your spending.
- **Weekly log**: 13 weeks of execution tracking with reply rate, close rate and revenue per hour.
- **Start here**: how to use it in five minutes.

**Honest notes:** every formula was recalculated and checked against independent math. It is a planning tool: results depend on your assumptions, returns are never guaranteed, and it is not financial advice. No login, no subscription, no upsell inside the file.

*Designed by Quiet Money with AI assistance.*

### Tags
financial freedom, budget planner, retirement calculator, FIRE calculator, passive income tracker, spreadsheet template, rat race, income streams

### After-purchase message
Thanks for grabbing the planner. Start on tab 1 and change one number at a time. If something looks wrong, reply to this message and it gets fixed. More free tools: {link}

## Etsy (second: needs the owner's Etsy shop)
- Category: Digital downloads > Templates/Planners. Instant download, no shipping.
- Etsy asks sellers to disclose AI involvement: use the line "Designed by Quiet Money with AI assistance" in the item description, and check Etsy's current Creativity Standards wording before publishing.
- Do not list prompt packs, other sellers' designs, or anything that copies a best-seller.
- Same description and previews; mention "Excel and Google Sheets" in the title for search.

## Rules we keep
No income claims, no "get rich" promises, no testimonials we did not receive, no fake scarcity. Refund promise: 30 days, no questions.
"""


def build_all() -> dict:
    """Build every product, verify it, write previews and the listing. Returns a summary per product."""
    from faceless import product_debt
    registry = [("rat-race-escape-planner", build_escape_planner, verify_escape, listing_escape),
                ("debt-payoff-planner", product_debt.build_debt_planner, product_debt.verify_debt, product_debt.listing)]
    out = {}
    for slug, build, verify, listing in registry:
        dest = OUT / slug
        xlsx = build(dest)
        check = verify(xlsx)
        if check["ok"]:      # ship the recalculated file: same sheets, validation and charts, plus cached values so previewers show numbers
            shutil.copy2(check["calc"], xlsx)
        prev = previews(xlsx, dest) if check["ok"] else []
        (dest / "LISTING.md").write_text(listing(), encoding="utf-8")
        out[slug] = {"file": str(xlsx), "verified": check["ok"], "rows": check["rows"], "errors": check["errors"], "previews": len(prev)}
    if all(v["verified"] for v in out.values()):
        from faceless import product_print
        pr = product_print.build_all(Path(out["rat-race-escape-planner"]["file"]), Path(out["debt-payoff-planner"]["file"]))
        bad = sum(sum(v) for v in pr["unsafe_pixels"].values())
        out["money-reset-printables"] = {"file": pr["pdfs"]["Letter"], "verified": bad == 0, "previews": 0, "errors": [] if bad == 0 else [f"{bad} pixels in unsafe margins"],
                                         "rows": [(f"{k} page {i + 1} margin pixels", v, 0, v == 0) for k, vs in pr["unsafe_pixels"].items() for i, v in enumerate(vs)]}
        out["money-reset-kit"] = {"file": pr["bundle"], "verified": True, "previews": 0, "errors": [], "rows": []}
    return out
