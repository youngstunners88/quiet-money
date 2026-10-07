"""Engine unit tests: pure logic only (no network, no ffmpeg), so they run anywhere in seconds."""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from faceless import analytics, config, moneymath
from faceless.gauntlet import check_package, check_script, score
from faceless.pipeline.captions import callout_text, chunk_words, wrap_hook
from faceless.pipeline.ideate import similarity
from faceless.pipeline.script import normalize, word_count
from faceless.pipeline.voice import align_beats, attach_display
from faceless.providers.tts import rate_to_speed, speed_to_rate
from faceless.state import Job, TransitionError

SEED = Path(__file__).resolve().parent.parent / "script-lab" / "final" / "_seed-001-ronald-read.json"


def words(*tokens, step=0.3):
    return [{"w": t, "start": round(i * step, 3), "end": round(i * step + step * 0.9, 3)} for i, t in enumerate(tokens)]


# ---- money math -------------------------------------------------------------------------------

def test_future_value_matches_known_value():
    assert round(moneymath.future_value_monthly(200, 0.08, 40)) == 698202


def test_card_payoff_with_fixed_payment():
    months, interest = moneymath.card_payoff(3000, 0.24, 300)
    assert months == 12 and 370 < interest < 400


def test_card_payoff_never_ends_when_payment_below_interest():
    months, interest = moneymath.card_payoff(10000, 0.24, 100)
    assert interest == float("inf")


def test_fact_sheet_mentions_assumptions():
    sheet = moneymath.fact_sheet()
    assert "8%/yr" in sheet and "24% APR" in sheet


# ---- voice alignment ---------------------------------------------------------------------------

def test_attach_display_restores_punctuation():
    w = attach_display(words("Not", "his", "friends", "Not", "his", "neighbors"), "Not his friends. Not his neighbors.")
    assert [x["display"] for x in w] == ["Not", "his", "friends.", "Not", "his", "neighbors."]


def test_attach_display_merges_split_tokens():
    w = attach_display(words("J", "C", "Penney", "for"), "J.C. Penney for")
    assert [x["display"] for x in w] == ["J.C.", "Penney", "for"]
    assert w[0]["end"] == pytest.approx(0.57)


def test_attach_display_keeps_raw_words_when_streams_disagree():
    raw = words("eight", "million", "dollars")
    assert attach_display(raw, "$8M") == raw


def test_align_beats_starts_at_zero_and_is_contiguous():
    ws = words("One", "two", "three", "four", "five", "six")
    spans = align_beats(["One two three.", "Four five six."], ws, total=2.0)
    assert spans[0][0] == 0.0
    assert spans[0][1] == spans[1][0] == pytest.approx(0.9)
    assert spans[1][1] == 2.0


def test_rate_speed_roundtrip():
    assert rate_to_speed("+14%") == pytest.approx(1.14)
    assert speed_to_rate(0.94) == "-6%"


# ---- captions ----------------------------------------------------------------------------------

def test_chunks_break_at_sentence_end():
    ws = [{"w": t, "display": t, "start": 0, "end": 0} for t in "He pumped gas. Then he cleaned floors.".split()]
    chunks = [" ".join(x["display"] for x in c) for c in chunk_words(ws, 3, 16)]
    assert "gas." in chunks[0] and all("gas. Then" not in c for c in chunks)


def test_chunks_do_not_end_on_weak_word():
    ws = [{"w": t, "display": t, "start": 0, "end": 0} for t in "He bought a house in the city".split()]
    for c in chunk_words(ws, 3, 16)[:-1]:
        assert c[-1]["display"].lower() not in {"a", "the", "in"}


def test_hook_wraps_without_escaping_newlines():
    lines = wrap_hook("A janitor died with $8 million")
    assert len(lines) == 2 and all("\\" not in l for l in lines)


def test_long_callout_wraps_smaller():
    text, size = callout_text("25 years at a gas pump")
    assert "\\N" in text and size < 150
    assert callout_text("$698,000") == ("$698,000", 150)


# ---- scripts + gauntlet ------------------------------------------------------------------------

def seed_script():
    return normalize(json.loads(SEED.read_text(encoding="utf-8")))


def test_seed_script_passes_code_gates():
    gates = check_script(seed_script(), history=[])
    s, hard = score(gates)
    assert not hard, [g.name for g in hard]
    assert s >= 90


def test_compliance_gate_blocks_promises():
    scr = seed_script()
    scr["beats"][3]["say"] = "This fund has guaranteed returns, you should buy it now."
    hard = {g.name for g in check_script(scr, history=[]) if g.hard and not g.passed}
    assert "compliance_phrases" in hard


def test_originality_gate_blocks_duplicates():
    scr = seed_script()
    hard = {g.name for g in check_script(scr, history=[scr["title"] + " " + scr["beats"][0]["say"]])
            if g.hard and not g.passed}
    assert "originality" in hard


def test_judge_compliance_ignores_factual_score(monkeypatch):
    """Compliance risk 4 with stated assumptions passes; factual risk is judged by its own gate."""
    from faceless import gauntlet
    monkeypatch.setattr(gauntlet.llm, "complete", lambda *a, **k: {
        "hook": 8, "retention": 7, "value": 7, "factual_risk": 6, "compliance_risk": 4,
        "suspect_claims": ["x"], "fixes": []})
    monkeypatch.setattr("faceless.decide.RECORDS", Path("/dev/null"))
    monkeypatch.setattr("faceless.events.JOURNAL", Path("/dev/null"))
    gates = {g.name: g for g in gauntlet.judge_script(seed_script())[0]}
    assert gates["judge_compliance"].passed
    assert not gates["judge_facts"].passed


def test_word_count_hard_only_when_far_out_of_range():
    scr = seed_script()
    gate = lambda s: next(g for g in check_script(s, history=[]) if g.name == "word_count")  # noqa: E731
    near = dict(scr, beats=scr["beats"][:-1])            # a few words short: soft
    far = dict(scr, beats=scr["beats"][:8])              # ~60 words short: hard, rewrite in the script loop
    assert not gate(near).hard or gate(near).passed
    assert gate(far).hard and not gate(far).passed


def test_length_brief_gives_a_word_target():
    from faceless.orchestrator import length_brief
    brief = length_brief(53.3, 124)
    assert "too short" in brief and "about 187 words" in brief


def test_normalize_strips_hook_punctuation():
    assert normalize({"title": "t", "hook_text": "21 YEARS.", "beats": []})["hook_text"] == "21 YEARS"


def test_normalize_clears_callout_on_hook_beat():
    scr = normalize({"title": "t", "beats": [{"say": "a b", "callout": "X", "visual": "v"}]})
    assert scr["beats"][0]["callout"] == ""


def test_word_count():
    assert word_count({"beats": [{"say": "one two"}, {"say": "three"}]}) == 3


def test_package_gate_requires_disclaimers():
    bad = {"description": "great video", "caption": "#money"}
    good = {"description": "Educational content, not financial advice. Narration and visuals are AI-assisted.",
            "caption": "#money"}
    assert any(g.hard and not g.passed for g in check_package(bad))
    assert not any(g.hard and not g.passed for g in check_package(good))


def test_similarity_is_symmetric_and_bounded():
    a, b = "janitor died with 8 million", "the janitor who died with $8 million"
    assert 0 <= similarity(a, b) <= 1 and similarity(a, b) == similarity(b, a)


# ---- routing + state ---------------------------------------------------------------------------

def test_allocate_balanced_by_default():
    weights = {p.id: 1.0 for p in config.load()["pillars"]}
    plan = analytics.allocate(len(weights), weights, day_index=3)
    assert sorted(plan) == sorted(weights)


def test_allocate_favours_winners_but_caps_at_two():
    weights = {"story": 3.0, "math": 1.0, "psychology": 0.4, "myth": 0.4, "playbook": 0.4}
    plan = analytics.allocate(5, weights, day_index=0)
    assert plan.count("story") == 2 and len(plan) == 5


def test_job_cannot_move_backwards(tmp_path, monkeypatch):
    monkeypatch.setattr("faceless.config.Paths.jobs", tmp_path)
    monkeypatch.setattr("faceless.events.JOURNAL", tmp_path / "journal.jsonl")
    job = Job(id="t", pillar="story", topic="x", day="2026-01-01")
    job.advance("scripted")
    job.advance("voiced")
    with pytest.raises(TransitionError):
        job.advance("scripted")
    assert job.reached("scripted") and not job.reached("rendered")


def test_playbook_needs_three_steps():
    from faceless.gauntlet import check_structure
    beats = lambda *xs: [{"say": x} for x in xs]  # noqa: E731
    one = {"pillar": "playbook", "beats": beats("Hook.", "Step one. Open your bank app.", "Then save.")}
    three = {"pillar": "playbook", "beats": beats("Hook.", "Step one. Open the app.", "Step two. Set a transfer.",
                                                 "Step three. Turn on alerts.")}
    assert not check_structure(one, "")[0].passed
    assert check_structure(three, "")[0].passed
    assert check_structure({"pillar": "story", "beats": []}, "") == []


def test_spoken_words_expand_numbers():
    from faceless.pipeline.script import spoken_words
    assert spoken_words("$226,000") == 6          # two hundred twenty six thousand dollars
    assert spoken_words("in 2014") == 3           # in twenty fourteen
    assert spoken_words("8% a year") == 4         # eight percent a year
    assert spoken_words("plain words only") == 3


def test_normalize_caps_callouts_keeping_numbers():
    beats = [{"say": "x", "callout": c, "visual": "v"} for c in
             ["", "A", "B", "$1", "C", "$2", "D", "E", "$3", "F", "G"]]
    out = normalize({"title": "t", "beats": beats})["beats"]
    kept = [b["callout"] for b in out if b["callout"]]
    assert len(kept) == 7 and {"$1", "$2", "$3"} <= set(kept)


def test_composio_key_alias_and_execute_parsing(monkeypatch):
    import os
    from faceless.providers import composio_tools as ct
    monkeypatch.delenv("COMPOSIO_API_KEY", raising=False)
    monkeypatch.setenv("COMPOSIO_API", "ak_test_alias")
    ct._ensure_key()
    assert os.environ["COMPOSIO_API_KEY"] == "ak_test_alias"

    class FakeSession:
        def execute(self, slug, arguments):
            assert slug == "YOUTUBE_LIST_CHANNELS" and arguments == {"mine": True}
            return {"data": {"items": [{"id": "UC1"}]}, "successful": True, "error": None, "log_id": "log_123"}

    monkeypatch.setattr(ct, "session", lambda: FakeSession())
    monkeypatch.setattr("faceless.events.JOURNAL", Path("/dev/null"))
    monkeypatch.setattr("faceless.ledger.LEDGER", Path("/dev/null"))
    res = ct.execute("YOUTUBE_LIST_CHANNELS", {"mine": True})
    assert res["log_id"] == "log_123" and res["data"]["items"][0]["id"] == "UC1"


def test_composio_execute_raises_with_log_id(monkeypatch):
    from faceless.providers import ProviderError
    from faceless.providers import composio_tools as ct

    class FailSession:
        def execute(self, slug, arguments):
            return {"data": {}, "successful": False, "error": "token expired", "log_id": "log_bad"}

    monkeypatch.setattr(ct, "session", lambda: FailSession())
    monkeypatch.setattr("faceless.events.JOURNAL", Path("/dev/null"))
    monkeypatch.setattr("faceless.ledger.LEDGER", Path("/dev/null"))
    with pytest.raises(ProviderError, match="log_bad"):
        ct.execute("YOUTUBE_LIST_CHANNELS", {})


def test_voice_classifier_separates_slop_from_our_voice():
    from faceless import quality
    slop = ("In today's fast-paced world, managing money can feel overwhelming. Let's dive into simple yet "
            "powerful tips that will transform your financial journey and unlock the life you deserve.")
    ours = ("A $6,000 card at 24% takes about 21 years on minimum payments. You'd pay $10,887 in interest. "
            "Pay $300 a month and it's gone in 26 months. Same debt. Two different lives.")
    assert quality.classify(slop)[0] == "slop"
    assert quality.classify(ours)[0] == "quality"


def test_aeo_name_parsing_and_self_match():
    from faceless import aeo
    assert aeo._names('["Graham Stephan", "Quiet Money", "The Financial Diet"]')[1] == "Quiet Money"
    assert aeo._names("1. Ali Abdaal\n2. QuietMoneyRules") == ["Ali Abdaal", "QuietMoneyRules"]
    assert aeo._is_us("QuietMoneyRules") and not aeo._is_us("Money Guy Show")


# ---- scheduling: resume + bank-ahead ----------------------------------------------------------

def test_extra_slots_roll_into_later_days():
    from faceless.providers.publish import nominal_slot
    assert nominal_slot("2026-10-02", 0) == ("2026-10-02", 0)
    assert nominal_slot("2026-10-02", 6) == ("2026-10-03", 1)
    assert nominal_slot("2026-10-02", 12) == ("2026-10-04", 2)


def test_open_slots_resume_and_bank_ahead(monkeypatch):
    from datetime import datetime, timedelta, timezone

    from faceless import orchestrator
    today = datetime.now(timezone.utc).date()
    done = [Job(id=f"j{s}", pillar="math", topic="t", day=today.isoformat(), slot=s, status=st)
            for s, st in [(0, "published"), (1, "packaged"), (2, "held"), (5, "published")]]
    monkeypatch.setattr(orchestrator, "all_jobs", lambda: done)
    d0, d1 = today.isoformat(), (today + timedelta(days=1)).isoformat()
    dawn = datetime(today.year, today.month, today.day, 1, 0, tzinfo=timezone.utc)   # before every slot
    assert orchestrator.open_slots(5, now=dawn) == [(d0, 2), (d0, 3), (d0, 4)]        # held slot is retried
    assert orchestrator.open_slots(5, extra=5, now=dawn) == [(d0, 2), (d0, 3), (d0, 4), (d1, 1), (d1, 2)]
    # 23:00 UTC is 19:00 in New York: only the 21:00 slot is still ahead; passed slots are never refilled
    late = datetime(today.year, today.month, today.day, 23, 0, tzinfo=timezone.utc)
    assert orchestrator.open_slots(5, now=late) == [(d0, 4)]


def test_dotenv_loads_without_overriding(tmp_path, monkeypatch):
    f = tmp_path / ".env"
    f.write_text('# comment\nQM_TEST_A="from-file"\nexport QM_TEST_B=b\nQM_TEST_EMPTY=\n')
    monkeypatch.setenv("QM_TEST_B", "from-env")
    monkeypatch.delenv("QM_TEST_A", raising=False)
    assert config.load_dotenv(f) == 1
    import os
    assert os.environ["QM_TEST_A"] == "from-file" and os.environ["QM_TEST_B"] == "from-env"
    monkeypatch.delenv("QM_TEST_A")
    f.write_text("QM_TEST_C=abc123  # personal account\nQM_TEST_D='a # b'\n")
    monkeypatch.delenv("QM_TEST_C", raising=False)
    monkeypatch.delenv("QM_TEST_D", raising=False)
    config.load_dotenv(f)
    assert os.environ["QM_TEST_C"] == "abc123" and os.environ["QM_TEST_D"] == "a # b"
    monkeypatch.delenv("QM_TEST_C")
    monkeypatch.delenv("QM_TEST_D")


def test_paid_image_fallback_is_capped(tmp_path, monkeypatch):
    from faceless import ledger
    from faceless.providers import ProviderUnavailable, images
    monkeypatch.setenv("OPENROUTER_API_KEY", "test")
    monkeypatch.setattr(ledger, "used", lambda provider, unit, day=None: 60 if provider == "openrouter" else 0)
    with pytest.raises(ProviderUnavailable):
        images.openrouter("a quiet kitchen table", 864, 1536, 1, tmp_path / "x.jpg")


def test_allocate_rotates_the_skipped_pillar():
    w = {"story": 1.0, "math": 1.0, "psychology": 1.0, "myth": 1.0, "playbook": 1.0, "escape": 1.5}
    plans = [analytics.allocate(5, w, d) for d in range(6)]
    assert all("escape" in p and len(set(p)) == 5 for p in plans)
    skipped = {next(k for k in w if k not in p) for p in plans}
    assert skipped == {"story", "math", "psychology", "myth", "playbook"}


def test_passive_income_math():
    assert moneymath.capital_for_income(3000) == 900000
    assert moneymath.months_to_target(1000, 0.0, 12000) == 12


def test_escape_series_needs_the_catch_and_a_move():
    from faceless.gauntlet import check_structure
    beats = lambda *says: {"pillar": "escape", "beats": [{"say": s} for s in says]}  # noqa: E731
    ok = beats("Two people sell the same skill.", "One charges ten times more.", "Here's the value equation.",
               "The catch: it takes years of reps.", "This week, rewrite your offer around the outcome.")
    assert check_structure(ok, "")[0].passed
    assert not check_structure(beats("Two people.", "Same skill.", "Different price.", "Do it."), "")[0].passed


def test_research_brief_drops_sources_and_unknown_ids():
    from faceless import research
    assert "## Sources" not in research.brief("kiyosaki") and "Rich Dad Poor Dad" in research.brief("kiyosaki")
    assert research.brief("nobody") == "" and research.brief(None) == ""


def test_measured_duration_overrides_hard_word_count():
    from faceless.gauntlet import Gate, reconcile
    gates = [Gate("word_count", False, 8, True), Gate("duration", True, 10, True)]
    assert not reconcile(gates)[0].hard
    assert reconcile([Gate("word_count", False, 8, True), Gate("duration", False, 10, True)])[0].hard


def test_blocked_provider_is_skipped_for_the_day(tmp_path, monkeypatch):
    from faceless import ledger
    from faceless.providers import ProviderUnavailable, images
    monkeypatch.setenv("CLOUDFLARE_API_KEY", "t")
    monkeypatch.setenv("CLOUDFLARE_ACCOUNT_ID", "a")
    monkeypatch.setattr(ledger, "used", lambda provider, unit, day=None: 1 if unit == "blocked" else 0)
    with pytest.raises(ProviderUnavailable):
        images.cloudflare("a desk", 864, 1536, 1, tmp_path / "x.jpg")



# ---- audit fixes ------------------------------------------------------------------------------

def test_redact_scrubs_query_keys_and_secret_env_values(monkeypatch):
    monkeypatch.setenv("QM_FAKE_API_KEY", "AIzaFAKEFAKEFAKEFAKE12345")
    msg = "Max retries exceeded with url: /v1beta/x:generateContent?key=AIzaFAKEFAKEFAKEFAKE12345 (Caused by...)"
    out = config.redact(msg + " again AIzaFAKEFAKEFAKEFAKE12345")
    assert "AIzaFAKE" not in out and "[redacted]" in out


def test_journal_never_stores_a_key(tmp_path, monkeypatch):
    from faceless import events
    monkeypatch.setattr(events, "JOURNAL", tmp_path / "j.jsonl")
    monkeypatch.setenv("QM_FAKE_TOKEN", "tok_FAKEFAKEFAKEFAKE")
    events.emit("PROVIDER_FAIL", error="401 for https://x.test/v1?token=tok_FAKEFAKEFAKEFAKE")
    assert "tok_FAKE" not in (tmp_path / "j.jsonl").read_text()


def test_gemini_tts_sends_key_in_header(tmp_path, monkeypatch):
    from faceless.providers import tts
    seen = {}

    class Boom:
        def post(self, url, **kw):
            seen.update(kw, url=url)
            raise tts.requests.ConnectionError(f"Max retries exceeded with url: {url}")
    monkeypatch.setenv("GEMINI_API_KEY", "AIzaFAKEFAKEFAKEFAKE12345")
    monkeypatch.setattr(tts, "http", lambda: Boom())
    with pytest.raises(tts.ProviderError) as err:
        tts.gemini("hello", tmp_path / "v.wav")
    assert seen["headers"]["x-goog-api-key"].startswith("AIza") and "params" not in seen
    assert "AIza" not in str(err.value) and "AIza" not in seen["url"]


def test_json_ld_cannot_break_out_of_its_script_block():
    from faceless.site import ld_json
    out = ld_json({"headline": "x</script><script>alert(1)</script> & more"})
    assert "</script>" not in out and "<" not in out and "&" not in out
    assert json.loads(out)["headline"].startswith("x</script>")


def test_escape_gate_needs_whole_words():
    from faceless.gauntlet import check_structure
    beats = {"pillar": "escape", "beats": [{"say": s} for s in (
        "A realistic number.", "Listen closely.", "Honestly, it's brisk work.", "A specialist knows.", "That's it.")]}
    assert not check_structure(beats, "")[0].passed


def test_paid_image_cap_counts_requests_in_flight(tmp_path, monkeypatch):
    from faceless import ledger
    from faceless.providers import ProviderUnavailable, images
    monkeypatch.setenv("OPENROUTER_API_KEY", "test")
    monkeypatch.setattr(ledger, "used", lambda provider, unit, day=None: 59 if unit == "images" else 0)
    monkeypatch.setattr(images, "_or_inflight", 1)   # another worker already holds the 60th
    with pytest.raises(ProviderUnavailable):
        images.openrouter("a desk", 864, 1536, 1, tmp_path / "x.jpg")


def test_failed_slot_never_fails_the_previous_finished_job(tmp_path, monkeypatch):
    from faceless import events, orchestrator
    from faceless.state import Job
    monkeypatch.setattr("faceless.config.Paths.jobs", tmp_path)
    monkeypatch.setattr(events, "JOURNAL", tmp_path / "j.jsonl")
    monkeypatch.setattr(orchestrator, "open_slots", lambda count, extra=0: [("2026-10-02", 0), ("2026-10-02", 1)])
    monkeypatch.setattr(orchestrator, "all_jobs", lambda: [])
    monkeypatch.setattr(orchestrator, "write_daily_report", lambda jobs: None)
    monkeypatch.setattr(orchestrator.Paths, "ensure", lambda: None)
    calls = iter([{"topic": "first"}, RuntimeError("no fresh topics")])

    def next_topic(pillar, reserved):
        item = next(calls)
        if isinstance(item, Exception):
            raise item
        return item
    monkeypatch.setattr(orchestrator.ideate, "next_topic", next_topic)
    monkeypatch.setattr(orchestrator, "start_job",
                        lambda pillar, item, slot: Job(id=f"j{slot}", pillar=pillar, topic=item["topic"], day="2026-10-02", slot=slot))

    def produce(job, **kw):
        job.status = "packaged"
    monkeypatch.setattr(orchestrator, "produce", produce)
    jobs = orchestrator.daily(5)


# ---- audit regressions (2026-10-02) -----------------------------------------------------------

def test_journal_scrubs_secrets(tmp_path, monkeypatch):
    from faceless import events
    monkeypatch.setattr(events, "JOURNAL", tmp_path / "journal.jsonl")
    monkeypatch.setenv("QM_TEST_API_KEY", "AIzaSyTESTSECRETVALUE123")
    events.emit("PROVIDER_FAIL", error="Max retries exceeded with url: /v1/m:generateContent?key=AIzaOTHER123&alt=json "
                                       "token AIzaSyTESTSECRETVALUE123")
    line = (tmp_path / "journal.jsonl").read_text()
    assert "AIzaSyTESTSECRETVALUE123" not in line and "AIzaOTHER123" not in line and "alt=json" in line


def test_gemini_key_never_in_url(monkeypatch):
    from faceless.providers import llm
    seen = {}

    class FakeSession:
        def post(self, url, **kw):
            seen.update(url=url, **kw)
            raise llm.requests.ConnectionError("down")

    monkeypatch.setenv("GEMINI_API_KEY", "secret-gemini-key")
    monkeypatch.setattr(llm, "http", lambda: FakeSession())
    with pytest.raises(llm.ProviderError):
        llm.gemini("hi", system=None, want_json=False, temperature=0.5)
    assert "params" not in seen and "secret-gemini-key" not in seen["url"]
    assert seen["headers"]["x-goog-api-key"] == "secret-gemini-key"


def test_site_jsonld_cannot_close_its_script_tag():
    from faceless import site
    out = site.ld_json({"headline": "</script><script>alert(1)</script> & more"})
    assert "</script>" not in out and "<" not in out and json.loads(out)["headline"].startswith("</script>")


def test_site_tolerates_string_facts_and_keeps_acronyms():
    from faceless import site
    assert site._items(["a claim", {"claim": "b"}, None, 3], ("claim", "basis")) == [
        {"claim": "a claim", "basis": ""}, {"claim": "b"}, {"claim": "3", "basis": ""}]
    assert site.sentence_case("HE STARTED WITH $820") == "He started with $820"
    assert site.sentence_case("the IRA trick") == "The IRA trick"
    assert site.rfc822("2026-10-02").startswith("Fri, 02 Oct 2026") and site.rfc822("bad") == ""


def test_judge_zero_risk_is_not_read_as_five(monkeypatch):
    from faceless import gauntlet
    monkeypatch.setattr(gauntlet.llm, "complete", lambda *a, **k: {
        "hook": 8, "retention": 7, "value": 7, "factual_risk": 0, "compliance_risk": 0,
        "fixes": "tighten beat 3", "suspect_claims": None})
    monkeypatch.setattr("faceless.decide.RECORDS", Path("/dev/null"))
    monkeypatch.setattr("faceless.events.JOURNAL", Path("/dev/null"))
    gates, _ = gauntlet.judge_script(json.loads(SEED.read_text(encoding="utf-8")))
    by = {g.name: g for g in gates}
    assert by["judge_facts"].passed and by["judge_compliance"].passed
    assert gauntlet.judge_number("7/10") == 7 and gauntlet.judge_number(None) == 5


def test_judge_non_object_answer_is_advisory(monkeypatch):
    from faceless import gauntlet
    monkeypatch.setattr(gauntlet.llm, "complete", lambda *a, **k: ["not", "an", "object"])
    monkeypatch.setattr("faceless.events.JOURNAL", Path("/dev/null"))
    assert gauntlet.judge_script({"title": "t", "beats": [{"say": "x"}]}) == ([], {})


def test_bad_metrics_and_backlog_lines_are_skipped(tmp_path, monkeypatch):
    from faceless.pipeline import ideate
    m = tmp_path / "metrics.jsonl"
    m.write_text('{"pillar": "story", "views": 10}\nnot json\n[1]\n{"pillar": "math", "views": null}\n')
    monkeypatch.setattr(analytics, "METRICS", m)
    assert len(analytics.load()) == 2
    b = tmp_path / "backlog.jsonl"
    b.write_text('{"pillar": "story", "topic": "a"}\n{"pillar": "story"\n{"topic": "no pillar"}\n')
    monkeypatch.setattr(ideate, "BACKLOG", b)
    monkeypatch.setattr("faceless.events.JOURNAL", tmp_path / "journal.jsonl")
    assert [i["topic"] for i in ideate.load_backlog()] == ["a"]


def test_normalize_coerces_llm_shapes():
    s = normalize({"title": "t", "hashtags": "#money #debt", "facts": ["x", {"claim": "c"}],
                   "beats": ["junk", {"say": "one"}, {"say": "two", "callout": "$5\nNOW"}]})
    assert s["hashtags"] == ["#money", "#debt"]
    assert s["facts"] == [{"claim": "x", "basis": ""}, {"claim": "c"}]
    assert [b["say"] for b in s["beats"]] == ["one", "two"] and s["beats"][1]["callout"] == "$5 NOW"


def test_caption_escape_strips_line_breaks_and_tags():
    from faceless.pipeline.captions import esc
    assert esc("a\n{\\b1}Dialogue") == "a (b1)Dialogue"


def test_topic_error_does_not_fail_previous_slot(monkeypatch):
    from faceless import orchestrator

    class FakeJob:
        def __init__(self, topic, slot):
            self.id, self.topic, self.slot, self.status, self.notes = f"j{slot}", topic, slot, "planned", []
            self.day = "2026-10-02"

        def advance(self, to, **_):
            self.status = to

        def save(self):
            pass

    topics = iter([{"topic": "first"}])

    def next_topic(pillar, reserved):
        item = next(topics, None)
        if item is None:
            raise RuntimeError("no fresh topics")
        return item

    monkeypatch.setattr(orchestrator, "open_slots", lambda count, extra=0: [("2026-10-02", 0), ("2026-10-02", 1)])
    monkeypatch.setattr(orchestrator.ideate, "next_topic", next_topic)
    monkeypatch.setattr(orchestrator, "start_job", lambda pillar, item, slot: FakeJob(item["topic"], slot))
    monkeypatch.setattr(orchestrator, "produce", lambda job, **kw: setattr(job, "status", "packaged"))
    monkeypatch.setattr(orchestrator, "all_jobs", lambda: [])
    monkeypatch.setattr(orchestrator, "write_daily_report", lambda jobs: None)
    monkeypatch.setattr(orchestrator.Paths, "ensure", classmethod(lambda cls: None))
    monkeypatch.setattr("faceless.events.JOURNAL", Path("/dev/null"))
    jobs = orchestrator.daily(2)
    assert [j.status for j in jobs] == ["packaged"]


def test_keyword_demand_counts_related_phrases_and_survives_outages(monkeypatch):
    from faceless import keywords
    monkeypatch.setattr(keywords, "suggest", lambda seed: ("quit job calculator", "quit job funny", "unrelated thing"))
    d = keywords.demand("Quit Your Job Forever")
    assert d["score"] == 2 and "unrelated thing" not in d["phrases"]
    monkeypatch.setattr(keywords, "suggest", lambda seed: ())
    assert keywords.demand("Quit Your Job Forever")["score"] == 0


def test_laya_backend_answers_are_validated_and_dry_run_by_default(tmp_path, monkeypatch):
    from faceless import decide, events
    monkeypatch.setattr(decide, "RECORDS", tmp_path / "d.jsonl")
    monkeypatch.setattr(events, "JOURNAL", tmp_path / "j.jsonl")
    monkeypatch.setenv("LAYA_URL", "http://laya.test:8000/")
    monkeypatch.delenv("TYPESAFE_API_KEY", raising=False)

    class Resp:
        status_code = 200
        def json(self):
            return {"answers": {"next_step": {"choice": "publish", "answer_confidence": 0.97},
                                "rogue": {"choice": "x", "answer_confidence": 0.99}}, "usage": {"input_tokens": 12}}

    class Http:
        def post(self, url, **kw):
            assert url == "http://laya.test:8000/v1/systemone" and kw["json"]["questions"]["next_step"]["type"] == "choice"
            return Resp()
    monkeypatch.setattr("faceless.providers.http", lambda: Http())
    q = {"next_step": decide.Q("choice", "ship?", fallback="hold", options={"publish": "ok", "hold": "no"})}
    d = decide.ask(q, {"score": 90})["next_step"]
    assert decide._backend() == "laya" and d.jev_value == "publish" and d.value == "hold"   # dry-run: logged, not acted on


# ---- opportunity scanner ----------------------------------------------------------------------

def _health(**kw):
    h = {"days": 7, "videos": 10, "held": 0, "hold_rate": 0.0, "top_failing_gates": [], "images": {"cloudflare": 50},
         "provider_blocked_days": [], "paid_usd": 0.0, "backlog_unused": {p: 20 for p in ("story", "math")},
         "metrics_rows": 50, "publish_mode": "upload_post", "pillars": ["story", "math"]}
    h.update(kw)
    return h


def test_scout_quiet_operation_has_no_proposals():
    from faceless import scout
    assert scout.proposals(_health()) == []


def test_scout_flags_quota_golive_metrics_backlog_and_holds():
    from faceless import scout
    h = _health(images={"pollinations": 60, "cloudflare": 20}, provider_blocked_days=["2026-10-02"], publish_mode="local",
                metrics_rows=2, backlog_unused={"story": 3, "math": 20}, videos=10, held=4, hold_rate=0.4,
                top_failing_gates=[("word_count", 3)])
    ids = {p["id"]: p for p in scout.proposals(h)}
    assert set(ids) == {"images-quota", "go-live", "metrics", "topup-story", "hold-rate"}
    assert ids["images-quota"]["autonomy"] == ids["go-live"]["autonomy"] == "owner" and ids["topup-story"]["autonomy"] == "auto"
    assert "word_count" in ids["hold-rate"]["why"]


def test_scout_act_only_tops_up_backlogs(tmp_path, monkeypatch):
    from faceless import scout
    monkeypatch.setattr(scout, "REPORT", tmp_path / "o.md")
    monkeypatch.setattr(scout, "OPPS", tmp_path / "o.jsonl")
    from faceless import empire
    monkeypatch.setattr(empire, "DIR", tmp_path)
    monkeypatch.setattr(empire, "RANKING", tmp_path / "RANKING.md")
    monkeypatch.setattr(scout, "health", lambda days=7: _health(publish_mode="local", backlog_unused={"story": 2, "math": 20}))
    called = []
    monkeypatch.setattr(scout, "topup", lambda pillar, n=6: called.append(pillar) or [{"topic": "t"}])
    res = scout.run(act=True, extra=[{"id": "x", "title": "Find a sponsor", "category": "money", "impact": 5, "effort": 5,
                                      "autonomy": "owner", "why": "w"}])
    assert called == ["story"] and res["done"] == ["topup-story: +1 topics"]
    assert "Find a sponsor" in (tmp_path / "o.md").read_text() and (tmp_path / "o.jsonl").exists()
    report = (tmp_path / "o.md").read_text()
    assert "## Portfolio" in report and "Owner unlocks the most" in report and (tmp_path / "RANKING.md").exists()
    assert [p["id"] for p in res["proposals"]][0] == "topup-story"    # impact 3 / effort 1 outranks 5 / 5


def test_scout_topup_skips_stale_and_unsearched_topics(monkeypatch):
    from faceless import scout
    from faceless.pipeline import ideate
    monkeypatch.setattr(scout, "_signals", lambda p: ["save money fast"])
    monkeypatch.setattr(ideate, "history_texts", lambda: [])
    monkeypatch.setattr(ideate, "load_backlog", lambda: [])
    saved = []
    monkeypatch.setattr(ideate, "append_backlog", lambda items: saved.extend(items))
    monkeypatch.setattr("faceless.providers.llm.complete", lambda *a, **k: {"ideas": [
        {"topic": "Cancel one subscription tonight", "angle": "a", "hook": "h"},
        {"topic": "Zxqv blorp frobnicate", "angle": "a", "hook": "h"}]})
    monkeypatch.setattr(scout.keywords, "demand", lambda t: {"score": 0 if "Zxqv" in t else 5, "phrases": []})
    got = scout.topup("playbook", n=5)
    assert [g["topic"] for g in got] == ["Cancel one subscription tonight"] and saved == got and got[0]["source"] == "scout"


def test_gemini_tries_next_model_when_json_is_cut_off(monkeypatch):
    from faceless.providers import llm
    replies = iter(['{"title": "cut off", "beats": [', '{"title": "ok"}'])

    class R:
        status_code = 200

        def __init__(self, text):
            self.text = text

        def json(self):
            return {"candidates": [{"content": {"parts": [{"text": self.text}]}, "finishReason": "MAX_TOKENS"}]}

    class FakeSession:
        def post(self, url, **kw):
            return R(next(replies))

    monkeypatch.setenv("GEMINI_API_KEY", "k")
    monkeypatch.setattr(llm, "http", lambda: FakeSession())
    monkeypatch.setattr(llm.ledger, "spend", lambda *a, **k: None)
    assert llm.gemini("p", system=None, want_json=True, temperature=0.5) == '{"title": "ok"}'


def test_mostly_placeholder_art_is_a_hard_failure():
    from faceless.gauntlet import check_visuals
    imgs = lambda prov: [{"beat": i, "provider": p} for i, p in enumerate(prov)]  # noqa: E731
    few = {g.name: g for g in check_visuals(imgs(["cloudflare"] * 10 + ["procedural"] * 2), 12)}
    most = {g.name: g for g in check_visuals(imgs(["procedural"] * 12), 12)}
    assert not few["no_placeholder_art"].passed and not few["no_placeholder_art"].hard
    assert most["no_placeholder_art"].hard


def test_dark_cinematic_first_frame_is_visible_but_black_is_not():
    from faceless.gauntlet import frame_visible
    black = bytes([4] * 7296)
    low_key = bytes([8] * 6000 + [70] * 1296)       # mean ~19 at worst, lit detail in the top 18%
    darker = bytes([5] * 6200 + [60] * 1096)        # mean ~13, still a visible subject
    assert not frame_visible(black) and frame_visible(low_key) and frame_visible(darker)
    assert not frame_visible(b"")


def test_small_lit_subject_on_black_is_visible():
    from faceless.gauntlet import frame_visible
    # 2026-10-05 math video: a lit hand and coin on black, mean 14, 90th percentile 35
    assert frame_visible(bytes([6] * 6400 + [35] * 500 + [165] * 396))
    assert not frame_visible(bytes([3] * 6900 + [200] * 396))   # black frame with only the hook box lit


def test_scout_topup_accepts_a_bare_idea_list(monkeypatch, tmp_path):
    from faceless import scout
    from faceless.pipeline import ideate
    monkeypatch.setattr("faceless.providers.llm.complete", lambda *a, **k: [{"topic": "Why rent feels cheaper"}, "junk"])
    monkeypatch.setattr(scout.keywords, "demand", lambda t: {"score": 3})
    monkeypatch.setattr(scout, "_signals", lambda p: [])
    monkeypatch.setattr(ideate, "BACKLOG", tmp_path / "backlog.jsonl")
    monkeypatch.setattr(ideate, "history_texts", lambda: [])
    monkeypatch.setattr("faceless.events.JOURNAL", tmp_path / "j.jsonl")
    assert [a["topic"] for a in scout.topup("myth", 2)] == ["Why rent feels cheaper"]


# ---- HyperFrames motion cards -------------------------------------------------------------------

def test_card_number_parsing_and_years_are_static():
    from faceless.pipeline import cards
    n = cards.parse_number("$698,202 total")
    assert n["value"] == 698202 and n["prefix"] == "$" and n["label"] == "TOTAL" and not n["static"]
    assert cards.parse_number("8%")["suffix"] == "%" and cards.parse_number("$1.5M")["decimals"] == 1
    assert cards.parse_number("2014")["static"] and cards.parse_number("FREE") is None


def test_card_pick_respects_mode_and_never_touches_the_hook(monkeypatch):
    from faceless.pipeline import cards
    monkeypatch.setattr(cards, "available", lambda: True)
    script = {"beats": [{"say": "Hook line here", "callout": "$5"}, {"say": "Then 8 percent a year", "callout": "8%"},
                        {"say": "No number here at all", "callout": ""}]}
    imgs = [{"provider": "cloudflare"}, {"provider": "cloudflare"}, {"provider": "procedural"}]
    base = config.load()["production"]
    for mode, pillar, expect in [("off", "math", set()), ("auto", "math", {2}), ("numbers", "math", {1, 2}),
                                 ("numbers", "story", {2})]:
        monkeypatch.setitem(base, "cards", {"mode": mode, "pillars": ["math", "escape"]})
        assert set(cards.pick(script, imgs, pillar)) == expect, (mode, pillar)
    monkeypatch.setitem(base, "cards", {"mode": "numbers", "pillars": ["math"]})
    assert set(cards.pick(script, None, "math")) == {1}               # before images exist: numeric beats only
    monkeypatch.setattr(cards, "available", lambda: False)
    assert cards.pick(script, imgs, "math") == {}                      # no Node/Chrome: cards quietly off


def test_card_project_escapes_llm_text_and_times_scenes(tmp_path):
    from faceless.pipeline import cards
    specs = [(1, {"kind": "stat", "value": 1500.0, "prefix": "$", "suffix": "", "decimals": 0, "label": "<SCRIPT>X</SCRIPT>"}, 3.0),
             (2, {"kind": "phrase", "text": "<IMG SRC=X> HELLO & BYE"}, 2.0)]
    total = cards.build_project(tmp_path / "p", specs, "math", "#FFD23F", 30, kicker="<B>K</B>")
    page = (tmp_path / "p" / "index.html").read_text()
    assert total == 5.0 and "<SCRIPT>X" not in page and "<IMG SRC" not in page and "<B>K" not in page
    assert 'data-start="3.000"' in page and 'data-duration="2.000"' in page and "$1,500" in page
    assert page.count('class="num"') + page.count('class="num" style') >= 8      # stacked count-up values, not callbacks


def test_plan_shots_carries_beat_and_card_spec():
    from faceless.pipeline import render
    imgs = [{"path": "a.jpg"}, {"path": None, "card": {"kind": "phrase", "text": "X"}}]
    shots = render.plan_shots([(0.0, 3.0), (3.0, 6.0)], imgs, 30, 3.2, "seed")
    assert [s["beat"] for s in shots] == [0, 1] and shots[1]["card"] and shots[1]["image"] is None


def test_card_beats_skip_the_caption_callout(tmp_path):
    from faceless.pipeline import captions
    script = {"hook_text": "HOOK HERE NOW", "beats": [{"say": "one", "callout": ""}, {"say": "two", "callout": "$200"}]}
    voice = {"duration": 8.0, "beats": [[0.0, 3.0], [3.0, 8.0]], "words": [
        {"w": "one", "start": 0.1, "end": 0.5, "display": "one"}, {"w": "two", "start": 3.1, "end": 3.5, "display": "two"}]}
    a, b = tmp_path / "a.ass", tmp_path / "b.ass"
    captions.build(script, voice, a)
    captions.build(script, voice, b, skip_callouts={1})
    assert "Callout," in a.read_text().split("[Events]")[1] and "Callout," not in b.read_text().split("[Events]")[1]


def test_memory_indexes_scripts_and_filters_by_pillar_status(tmp_path, monkeypatch):
    from faceless import memory
    from faceless.config import Paths
    final, jobs, reports = tmp_path / "final", tmp_path / "jobs", tmp_path / "reports"
    for d in (final, jobs, reports):
        d.mkdir()
    monkeypatch.setattr(Paths, "final", final)
    monkeypatch.setattr(Paths, "jobs", jobs)
    monkeypatch.setattr(Paths, "reports", reports)
    mk = lambda job, pillar, title, say, status: (  # noqa: E731
        (final / f"{job}.json").write_text(json.dumps({"pillar": pillar, "title": title, "beats": [{"say": say}]})),
        (jobs / f"{job}.json").write_text(json.dumps({"status": status, "scores": {"gauntlet": 90}, "day": "2026-10-06", "notes": []})))
    mk("a", "math", "The minimum payment trap", "Paying only the minimum on a credit card takes 21 years", "published")
    mk("b", "myth", "Renting myth", "Renting is not throwing money away", "held")
    (reports / "b.json").write_text(json.dumps({"gates": [{"name": "word_count", "passed": False}, {"name": "hook", "passed": True}]}))
    db = tmp_path / "m.db"
    assert memory.build(db) == 2
    assert [r["job"] for r in memory.search("credit card minimum payments", db=db)] == ["a"]
    assert [r["job"] for r in memory.search("word_count", status="held", db=db)] == ["b"]
    assert memory.search("renting", pillar="math", db=db) == [] and memory.search('"; drop table docs; --', db=db) == []


# ---- empire gauntlet + forecast ----------------------------------------------------------------

def _opp(id_, stage="vetted", feeds=(), gates=(), build_gates=(), attack=None, **kw):
    base = {"id": id_, "name": id_, "rail": "test", "stage": stage, "feeds": list(feeds), "gates": list(gates),
            "build_gates": list(build_gates), "upside": 4, "leverage": 4, "automation": 4, "speed": 4, "effort": 2, "risk": 1,
            "kill": "no sales in 30 days", "note": "n", "attack": attack or {}}
    return {**base, **kw}


def test_empire_attack_kills_on_blocked_terms_or_ip_only():
    from faceless import empire
    assert empire.attack(_opp("a", attack={"tos": "block"}))["killed"]
    assert empire.attack(_opp("a", attack={"ip": "block"}))["killed"]
    soft = empire.attack(_opp("a", attack={"tos": "watch", "saturation": "high", "agent": "partly"}))
    assert not soft["killed"] and soft["mult"] == pytest.approx(0.7 * 0.8 * 0.8)
    assert empire.score(_opp("a", attack={"saturation": "high"})) < empire.score(_opp("a"))


def test_empire_lane_follows_dependencies_and_owner_gates():
    from faceless import empire
    items = [_opp("base", stage="scaled"), _opp("pilot", stage="pilot"), _opp("idea", stage="idea", feeds=["base"]),
             _opp("child", feeds=["base"], gates=["shop"]), _opp("late", feeds=["idea"]), _opp("acct", build_gates=["channels"]),
             _opp("bad", attack={"ip": "block"})]
    by = {o["id"]: o for o in items}
    by["shipped"] = _opp("shipped", stage="built", feeds=["base"], gates=["shop"])
    by["kid"] = _opp("kid", feeds=["shipped"])              # a built parent unblocks its children
    lanes = {i: empire.lane(o, by) for i, o in by.items()}
    assert lanes == {"base": "running", "pilot": "running", "idea": "validate", "child": "agent", "late": "waiting",
                     "acct": "owner", "bad": "dead", "shipped": "launch", "kid": "agent"}
    assert "Launch needs the owner" in empire.next_action(by["shipped"], by) and empire.next_action(by["shipped"], by).startswith("built")
    assert "Launch needs the owner" in empire.next_action(by["child"], by) and "waits for idea" == empire.next_action(by["late"], by)


def test_empire_asks_count_each_item_once_per_gate_and_skip_dead_or_scaled(tmp_path, monkeypatch):
    from faceless import empire
    p = tmp_path / "portfolio.jsonl"
    rows = [_opp("a", gates=["shop"]), _opp("b", gates=["shop", "email"]), _opp("both", gates=["channels"], build_gates=["channels"]),
            _opp("done", stage="scaled", gates=["shop"]), _opp("bad", gates=["shop"], attack={"tos": "block"})]
    p.write_text("\n".join(json.dumps(r) for r in rows) + "\nnot json\n", encoding="utf-8")
    monkeypatch.setattr(empire, "PORTFOLIO", p)
    monkeypatch.setattr(empire, "DIR", tmp_path)
    monkeypatch.setattr(empire, "RANKING", tmp_path / "RANKING.md")
    one = empire.score(rows[0])
    got = {text: (total, ids) for text, total, ids in empire.asks(empire.rank())}
    assert got[empire.GATES["shop"]] == (pytest.approx(round(2 * one, 1)), ["a", "b"])
    assert got[empire.GATES["channels"]] == (pytest.approx(round(one, 1)), ["both"])
    empire.write()
    text = (tmp_path / "RANKING.md").read_text(encoding="utf-8")
    assert "## Killed by the attack round" in text and "**bad" in text and "## What the owner unlocks" in text


def test_shipped_portfolio_is_consistent():
    from faceless import empire
    items = empire.load()
    ids = {o["id"] for o in items}
    assert len(ids) == len(items) >= 20
    for o in items:
        assert o["stage"] in empire.STAGES and set(o["feeds"]) <= ids and o["kill"]
        assert set(o["gates"]) | set(o["build_gates"]) <= set(empire.GATES), o["id"]
        assert all(1 <= o[k] <= 5 for k in ("upside", "leverage", "automation", "speed", "effort", "risk")), o["id"]
    assert any(empire.attack(o)["killed"] for o in items)   # the StarNet-style ideas stay dead


def test_forecast_ladder_is_ordered_and_uses_assumptions(monkeypatch):
    from faceless import forecast
    monkeypatch.setattr(forecast, "measured", lambda: {"images_per_video": 12.0, "videos_measured": 3, "card_beats": {}})
    monkeypatch.setattr(forecast, "neurons_per_image", lambda: 417.0)
    low, base, high = forecast.ladder()
    assert low["views"] < base["views"] < high["views"] and low["total"] < base["total"] < high["total"]
    assert base["net"] == pytest.approx(base["total"] - base["cost"], abs=0.01)
    mc = forecast.monthly_cost(5)
    assert mc["images_per_month"] == 1800 and mc["cloudflare_paid"] > forecast.CF_PAID_BASE_MONTH
    assert forecast.monthly_cost(1)["openrouter"] < mc["openrouter"]


# ---- repurposing kit ---------------------------------------------------------------------------

KIT_SCRIPT = {
    "title": "The House Myth", "hook_text": "YOUR HOUSE IS NOT AN INVESTMENT", "pillar": "myth", "hashtags": ["#a", "#b", "#c", "#d"],
    "first_comment": "Rent or buy?", "description": "Why a home is not an asset by default.",
    "beats": [{"say": "Your house is not an investment.", "callout": ""},
              {"say": "Taxes and repairs eat returns. Look at the math...", "callout": "TAXES"},
              {"say": "Investing $500 a month for 30 years grows to $745,180.", "callout": "$745,180"},
              {"say": "Buy a home for shelter. This week, list every cost of owning yours.", "callout": ""},
              {"say": "Follow for the money rules school never taught you, because your house is...", "callout": ""}]}


def test_kit_text_drops_fragments_and_keeps_disclosure():
    from faceless import repurpose
    assert repurpose.sentences("One. Two... Three") == ["One.", "Three"]
    assert all(not s.endswith("...") for b in KIT_SCRIPT["beats"] for s in repurpose.sentences(b["say"]))
    th, li = repurpose.thread(KIT_SCRIPT), repurpose.linkedin(KIT_SCRIPT)
    assert th.startswith("1/") and repurpose.DISCLOSE in th and repurpose.DISCLOSE in li and "Follow for" not in th + li
    assert all(len(t.split(" ", 1)[1].split("\n\n")[0]) <= 280 for t in th.split("\n\n---\n\n"))
    assert li.count("#") <= 3 and "Look at the math" not in li


def test_kit_picks_the_biggest_number_and_the_action_sentence():
    from faceless import repurpose
    assert repurpose.the_number(KIT_SCRIPT).startswith("$745,180")
    assert repurpose.action_beat(KIT_SCRIPT) == "This week, list every cost of owning yours."
    item = repurpose.newsletter_item(KIT_SCRIPT, {"pillar": "myth"})
    assert "**The number**" in item and "**Try this:**" in item


def test_kit_builds_carousel_and_pin_without_a_cover(tmp_path):
    from faceless import repurpose
    n = repurpose.carousel(KIT_SCRIPT, None, tmp_path / "c")
    repurpose.pin(KIT_SCRIPT, {"description": "Why.\n\nMore"}, None, tmp_path)
    assert n == 5 == len(list((tmp_path / "c").glob("*.png")))   # hook + 3 body beats + CTA
    from PIL import Image
    assert Image.open(tmp_path / "c" / "01.png").size == repurpose.CAROUSEL and Image.open(tmp_path / "pin.png").size == repurpose.PIN
    assert "Educational" in (tmp_path / "pin.txt").read_text()


def test_weekly_issue_marks_unsendable_without_a_postal_address(tmp_path, monkeypatch):
    from faceless import repurpose
    from faceless.config import Paths
    kit = tmp_path / "queue" / "2026-10-07" / "slot1-myth" / "kit"
    kit.mkdir(parents=True)
    (kit / "newsletter.md").write_text("### item\n", encoding="utf-8")
    monkeypatch.setattr(Paths, "queue", tmp_path / "queue")
    monkeypatch.setattr(Paths, "root", tmp_path)
    issue = repurpose.weekly_issue(end="2026-10-08")
    assert issue and "### item" in issue.read_text() and "NOT SENDABLE" in issue.read_text()
    assert repurpose.weekly_issue(end="2026-11-30") is None


# ---- Flow lane ---------------------------------------------------------------------------------

def _flow_dirs(tmp_path, monkeypatch):
    from faceless import flow
    for name, sub in (("DIR", ""), ("INBOX", "inbox"), ("USED", "used"), ("LISTS", "lists")):
        monkeypatch.setattr(flow, name, tmp_path / sub if sub else tmp_path)
    (tmp_path / "inbox").mkdir()
    return flow


def test_flow_shotlist_is_three_distinct_safe_prompts_and_rotates_by_day(tmp_path, monkeypatch):
    flow = _flow_dirs(tmp_path, monkeypatch)
    a = flow.shotlist("2026-10-07").read_text()
    b = flow.shotlist("2026-10-08").read_text()
    assert a.count("## ") == 3 and a != b
    assert a.count("no people's faces") == 3 and a.count("no text or numbers") == 3 and "nothing here logs in" in a
    assert a == flow.shotlist("2026-10-07").read_text()      # deterministic


def test_flow_claim_needs_enabled_and_a_clip_long_enough(tmp_path, monkeypatch):
    import shutil as sh
    import subprocess as sp
    if not sh.which("ffmpeg") or not sh.which("ffprobe"):
        pytest.skip("ffmpeg not installed")
    flow = _flow_dirs(tmp_path, monkeypatch)
    monkeypatch.setattr(flow, "enabled", lambda: True)
    mk = lambda name, secs: sp.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i",   # noqa: E731
                                    "testsrc2=size=640x1136:rate=12", "-t", str(secs), "-pix_fmt", "yuv420p",
                                    str(tmp_path / "inbox" / name)], check=True)
    mk("any-short.mp4", 2)
    mk("math-long.mp4", 5)
    (tmp_path / "inbox" / "broken.mp4").write_bytes(b"not a video")
    job = Job(id="j1", pillar="story", topic="t", day="2026-10-07")
    monkeypatch.setattr(Job, "dir", property(lambda self: tmp_path / "job"))
    got = flow.claim(job, 3.0)
    assert got and got["source"] == "math-long.mp4" and Path(got["clip"]).exists() and Path(got["frame"]).exists()
    assert not (tmp_path / "inbox" / "math-long.mp4").exists() and (tmp_path / "used" / "j1-math-long.mp4").exists()
    assert flow.claim(job, 3.0) is None                       # the short clip and the broken file are both refused
    monkeypatch.setattr(flow, "enabled", lambda: False)
    mk("story-x.mp4", 5)
    assert flow.claim(job, 3.0) is None


def test_card_compare_and_steps_parse_and_apply_to_every_series(monkeypatch):
    from faceless.pipeline import cards
    assert cards.parse_compare("$698,202 VS $298,072")["numeric"] and not cards.parse_compare("ASSET VS LIABILITY")["numeric"]
    assert cards.parse_compare("2017 vs 2018") is None and cards.parse_compare("$5 VS FREE") is None       # years / mixed sides
    assert cards.parse_compare("A VS B VS C") is None and cards.parse_compare("") is None
    st = cards.parse_step("STEP TWO", "Step two, set up an automatic transfer on payday. Then relax.")
    assert st["n"] == 2 and st["gist"] == "set up an automatic transfer on payday" and cards.parse_step("5 STEPS") is None
    monkeypatch.setattr(cards, "available", lambda: True)
    script = {"beats": [{"say": "hook", "callout": "A VS B"}, {"say": "x", "callout": "MAX VS MIN"},
                        {"say": "Step one, open the app.", "callout": "STEP ONE"}, {"say": "plain", "callout": ""}]}
    base = config.load()["production"]
    monkeypatch.setitem(base, "cards", {"mode": "numbers", "pillars": ["math"]})
    got = cards.pick(script, None, "story")                  # a series outside the numbers list still gets variety cards
    assert {i: v["kind"] for i, v in got.items()} == {1: "compare", 2: "steps"}     # never the hook beat
    monkeypatch.setitem(base, "cards", {"mode": "numbers", "pillars": ["math"], "variety": False})
    assert cards.pick(script, None, "story") == {}


def test_card_project_draws_compare_and_steps_with_seekable_tweens(tmp_path):
    from faceless.pipeline import cards
    specs = [(1, cards.parse_compare("$10 VS $5"), 3.0), (2, cards.parse_compare("<b>X</b> VS Y"), 2.0),
             (3, cards.parse_step("STEP THREE", "Step three, <script>alert(1)</script> go."), 2.5)]
    total = cards.build_project(tmp_path, specs, "playbook", "#FFD23F", 30, kicker="PLAYBOOK")
    page = (tmp_path / "index.html").read_text(encoding="utf-8")
    assert total == pytest.approx(7.5) and "<script>alert" not in page and "&lt;B&gt;X&lt;/B&gt;" in page
    import re
    assert all(re.match(r'tl\.to\("#s\d+",\{opacity:0,duration:0\.001\}', m) for m in re.findall(r'tl\.to\(.{0,40}', page))   # only the fade-out


# ---- product factory ---------------------------------------------------------------------------

def test_escape_planner_expected_math_behaves():
    from faceless import products
    base = products.expected_escape()
    assert base["freedom"] == 900_000 and 35 < base["years"] < 37 and base["saved"] > 1
    assert products.expected_escape({"monthly": 1500})["years"] < base["years"]
    assert products.expected_escape({"invested": 2_000_000})["months"] == 0.0
    d = products.DEFAULTS     # the contribution leg of the year-20 balance matches the engine used by the videos
    lump = d["invested"] * (1 + base["real"] / 12) ** 240
    assert base["balance_year20"] - lump == pytest.approx(moneymath.future_value_monthly(d["monthly"], base["real"], 20))


def test_escape_planner_recalculates_to_the_same_numbers_in_a_real_spreadsheet_engine(tmp_path):
    from faceless import products
    xlsx = products.build_escape_planner(tmp_path)
    if products.recalc(xlsx) is None:
        pytest.skip("LibreOffice Calc not available")
    res = products.verify_escape(xlsx)
    assert res["errors"] == [] and res["ok"], res["rows"]


def test_escape_planner_has_no_hardcoded_results_and_the_disclaimer(tmp_path):
    from faceless import products
    from openpyxl import load_workbook
    wb = load_workbook(products.build_escape_planner(tmp_path))
    n = wb["1 Your numbers"]
    assert all(str(n[f"B{r}"].value).startswith("=") for r in range(16, 26))       # every result is a formula
    assert "not financial advice" in wb["Start here"]["A19"].value
    assert wb.sheetnames == products.SHEETS


# ---- autoresponder ----------------------------------------------------------------------------

def test_autoresponder_numbers_come_from_moneymath_and_pass_its_own_rules():
    from faceless import autoresponder, moneymath
    es = autoresponder.emails()
    assert [e["day"] for e in es] == list(range(8)) and autoresponder.problems(es) == []
    text = " ".join(e["body"] for e in es)
    n = autoresponder.numbers()
    assert n["sub10"] == "$1,800" and n["needs"] == "$2,000" and n["half_raise"] == "$125"
    assert moneymath.money(moneymath.future_value_monthly(25, 0.08, 30)) in text            # day 2, computed not typed
    assert n["min_interest"] in text and n["min_years"] in text
    assert all("{{unsubscribe_url}}" in e["footer"] and "{{postal_address}}" in e["footer"] for e in es)


def test_autoresponder_rules_catch_hype_length_and_missing_footer():
    from faceless import autoresponder
    bad = {"key": "x", "day": 1, "subject": "S" * 61, "preview": "p", "body": "Guaranteed returns! " + "word " * 200, "footer": "none", "sign": "q"}
    found = " | ".join(autoresponder.problems([bad]))
    assert "subject over 60" in found and "words (limit 190)" in found and "banned phrase" in found and "footer lacks" in found


def test_autoresponder_day7_mentions_the_planner_only_when_it_is_live(monkeypatch):
    from faceless import autoresponder
    assert "Escape Planner" not in autoresponder.emails()[-1]["body"]
    cfg = config.load()
    monkeypatch.setitem(cfg, "newsletter", {"planner_url": "https://example.com/p"})
    assert "https://example.com/p" in autoresponder.emails()[-1]["body"]


def test_debt_simulation_basics():
    from faceless import product_debt as pd
    debts = [("a", 1000, 0.24, 50), ("b", 500, 0.10, 25)]
    av, sn = pd.simulate(debts, 100, "Avalanche"), pd.simulate(debts, 100, "Snowball")
    assert av["plan_interest"] <= sn["plan_interest"] + 1e-9            # highest rate first never costs more interest
    assert av["plan_months"] < av["base_months"] and av["plan_interest"] < av["base_interest"]
    assert pd.simulate(debts, 0)["plan_months"] <= pd.simulate(debts, 0)["base_months"]
    assert pd.simulate([("z", 0, 0.2, 10)], 50)["plan_months"] == 0     # nothing to pay
    assert pd.expected_growth(0, 100, 0.08, 30)["value"] == pytest.approx(moneymath.future_value_monthly(100, 0.08, 30))


def test_debt_planner_recalculates_to_the_python_simulation_for_both_methods(tmp_path):
    from faceless import product_debt as pd
    from faceless import products
    first = pd.build_debt_planner(tmp_path / "a")
    if products.recalc(first) is None:
        pytest.skip("LibreOffice Calc not available")
    res = pd.verify_debt(first)
    assert res["errors"] == [] and res["ok"], res["rows"]
    over = {"method": "Snowball", "extra": 75, "debts": [("Card", 2500, 0.2199, 80), ("", None, None, None), ("Zero", 0, 0.1, 0),
                                                         ("Loan", 9000, 0.06, 180), ("Small", 400, 0.29, 25)]}
    second = pd.build_debt_planner(tmp_path / "b", **over)
    res2 = pd.verify_debt(second, **over)
    assert res2["errors"] == [] and res2["ok"], res2["rows"]


def test_debt_planner_flags_a_minimum_below_the_interest(tmp_path):
    from faceless import product_debt as pd
    from openpyxl import load_workbook
    wb = load_workbook(pd.build_debt_planner(tmp_path, debts=[("Bad", 10000, 0.30, 100)]))
    f = wb["1 Your debts"]["F4"].value
    assert "below the monthly interest" in f and "N(D4)<N(B4)*N(C4)/12" in f


# ---- brand kit service -------------------------------------------------------------------------

def test_brandkit_names_are_validated_and_svg_is_escaped():
    from faceless import brandkit
    assert brandkit.clean("  Calm   Orbit ") == "Calm Orbit" and brandkit.monogram("Calm Orbit") == "CO" and brandkit.monogram("Zed") == "Z"
    for bad in ("", "x" * 25, "<script>", "Quote\"Name", "Émile"):
        with pytest.raises(ValueError):
            brandkit.clean(bad)
    svg = brandkit.logo_svg("R&D 'Labs'", brandkit.PALETTES["gold"])
    assert "R&amp;D" in svg and "&#x27;Labs&#x27;" in svg and "<script" not in svg


def test_brandkit_builds_every_file_at_exact_sizes_with_text_inside_the_safe_area(tmp_path):
    from PIL import Image
    from faceless import brandkit
    for name, tag, pal in (("Calm Orbit", "Slow living for fast minds", "mint"), ("Iron Fork", "Cook it right, once and for all, every time", "coral"),
                           ("A", "", "gold"), ("The Very Long Channel Nam", "x" * 60, "sky")):
        if len(name) > 24:
            name = name[:24]
        r = brandkit.build(name, tag, pal, "niche", ["a", "b", "c", "d", "e"], out=tmp_path / name.replace(" ", "-"))
        d = Path(r["dir"])
        assert Image.open(d / "avatar.png").size == (800, 800) and Image.open(d / "banner-youtube.png").size == (2560, 1440)
        assert Image.open(d / "og-image.png").size == (1200, 630) and Image.open(d / "watermark.png").size == (150, 150)
        assert {"BRAND.md", "logo.svg", "favicon-512.png", "favicon-180.png", "favicon-32.png", "palette.png", "preview.png"} <= set(r["files"])
        x0, y0, x1, y1 = r["text_box"]
        left, top = (2560 - brandkit.SAFE_W) // 2, (1440 - brandkit.SAFE_H) // 2
        assert x0 >= left and x1 <= left + brandkit.SAFE_W and y0 >= top and y1 <= top + brandkit.SAFE_H, (name, r["text_box"])
        import zipfile
        assert "preview.png" not in zipfile.ZipFile(r["zip"]).namelist()
    with pytest.raises(ValueError):
        brandkit.build("Ok Name", palette="neon", out=tmp_path / "x")
